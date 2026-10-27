"""Offline tests: generate the messy sample files, clean them, compare with the known truth.
Run with:  python -m unittest -v
"""
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

import clean
import make_samples


class EndToEndTest(unittest.TestCase):
    def test_merge_matches_ground_truth(self):
        truth = {}
        for prefix, n, s in (("N", 40, 42), ("S", 35, 43), ("W", 30, 44)):
            for o in make_samples.make_orders(prefix, n, s):
                truth[o["id"]] = o
        del truth["S-1008"]                                  # the row with a missing quantity

        with tempfile.TemporaryDirectory() as d:
            files = make_samples.build(Path(d) / "in", seed=42)
            out = Path(d) / "report.xlsx"
            self.assertEqual(clean.main([*map(str, files), "-o", str(out)]), 0)
            wb = load_workbook(out)
            self.assertEqual(wb.sheetnames, ["Summary", "Clean Data", "Pivot", "By Customer", "Charts", "Issues"])

            rows = list(wb["Clean Data"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(rows), len(truth))           # duplicates and bad rows removed, nothing lost
            for oid, odate, month, cust, prod, qty, price, rev, _src in rows:
                t = truth[oid]
                self.assertEqual(odate.date(), t["date"], oid)  # every date format parsed correctly
                self.assertEqual((cust, prod, qty), (t["customer"], t["product"], t["qty"]), oid)
                self.assertAlmostEqual(price, t["price"])
                self.assertAlmostEqual(rev, round(t["qty"] * t["price"], 2))

            pivot = list(wb["Pivot"].iter_rows(values_only=True))
            self.assertEqual(pivot[-1][0], "Total")
            self.assertAlmostEqual(pivot[-1][-1], round(sum(t["qty"] * t["price"] for t in truth.values()), 2), places=2)
            self.assertEqual(len(wb["Charts"]._charts), 3)
            issues = [r[2] for r in wb["Issues"].iter_rows(min_row=2, values_only=True)]
            self.assertEqual(sum("duplicate" in i for i in issues), 2)
            self.assertEqual(sum("missing quantity" in i for i in issues), 1)


class UnitTests(unittest.TestCase):
    def test_column_aliases(self):
        m = clean.map_columns(["Purchase Date", "OrderID", "SKU Name", "Units", "Price (USD)", "Coupon"])
        self.assertEqual(set(m.values()), {"order_date", "order_id", "product", "quantity", "unit_price"})

    def test_day_first_detection_and_parsing(self):
        self.assertTrue(clean.detect_day_first(["03/04/2026", "25/04/2026"]))
        self.assertFalse(clean.detect_day_first(["03/04/2026", "04/25/2026"]))
        self.assertTrue(clean.detect_day_first(["03.04.2026"]))
        self.assertEqual(clean.parse_date("03/04/2026", True), date(2026, 4, 3))
        self.assertEqual(clean.parse_date("03/04/2026", False), date(2026, 3, 4))
        self.assertEqual(clean.parse_date("2026-02-01", False), date(2026, 2, 1))
        self.assertEqual(clean.parse_date(46054, False), date(2026, 2, 1))   # Excel serial

    def test_numbers(self):
        self.assertEqual(clean.to_number("$1,234.50"), 1234.5)
        self.assertEqual(clean.to_number("1.234,50"), 1234.5)
        self.assertEqual(clean.to_number(" 12,50 "), 12.5)
        self.assertIsNone(clean.to_number(""))

    def test_canonical_spelling(self):
        import pandas as pd
        s = pd.Series(["ACME CORP", "Acme Corp", "acme corp", "Acme Corp"])
        self.assertEqual(set(clean.canonical_spelling(s)), {"Acme Corp"})


if __name__ == "__main__":
    unittest.main()
