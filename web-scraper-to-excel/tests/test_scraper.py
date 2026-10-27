"""Offline tests: run with  python -m unittest  (no network needed)."""
import tempfile
import unittest
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import load_workbook

import scraper

HTML = """
<ol>
<article class="product_pod"><img src="../media/a.jpg"><p class="star-rating Three"></p>
<h3><a href="a_1/index.html" title="Book  A">Book A</a></h3>
<p class="price_color">Â£51.77</p><p class="instock availability"> In stock </p></article>
<article class="product_pod"><img src="../media/a.jpg"><p class="star-rating Three"></p>
<h3><a href="a_1/index.html" title="Book A">Book A</a></h3>
<p class="price_color">£51.77</p><p class="availability">In stock</p></article>
<article class="product_pod"><img src="../media/b.jpg"><p class="star-rating Five"></p>
<h3><a href="b_2/index.html" title="Book B">Book B</a></h3>
<p class="price_color">£10.00</p><p class="availability">Out of stock</p></article>
</ol><ul class="pager"><li class="next"><a href="page-2.html">next</a></li></ul>
"""
PAGE = "https://books.toscrape.com/catalogue/page-1.html"


class ScraperTests(unittest.TestCase):
    def setUp(self):
        self.soup = BeautifulSoup(HTML, "html.parser")

    def test_parse_and_clean(self):
        rows = scraper.parse_page(self.soup, PAGE, 1)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0].title, "Book A")          # whitespace collapsed
        self.assertEqual(rows[0].price_gbp, 51.77)         # mojibake currency stripped
        self.assertEqual(rows[2].rating, 5)
        self.assertFalse(rows[2].in_stock)
        self.assertEqual(rows[0].product_url, "https://books.toscrape.com/catalogue/a_1/index.html")

    def test_next_page_and_dedupe(self):
        self.assertEqual(scraper.next_page_url(self.soup, PAGE), "https://books.toscrape.com/catalogue/page-2.html")
        rows = scraper.dedupe(scraper.parse_page(self.soup, PAGE, 1))
        self.assertEqual([r.title for r in rows], ["Book A", "Book B"])

    def test_export_formatting(self):
        rows = scraper.dedupe(scraper.parse_page(self.soup, PAGE, 1))
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "t.xlsx"
            scraper.export_xlsx(rows, out)
            ws = load_workbook(out)["Books"]
            self.assertTrue(ws["A1"].font.bold)
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertEqual(ws.max_row, 3)


if __name__ == "__main__":
    unittest.main()
