#!/usr/bin/env python3
"""Paginated web scraper that cleans, de-duplicates and exports to a formatted Excel file.

Default target: https://books.toscrape.com (a public sandbox site built for scraping practice).

Usage:
    python scraper.py                      # scrape all pages -> output/books.xlsx
    python scraper.py --max-pages 3        # first 3 pages only
    python scraper.py --out my.xlsx --delay 1.0
"""
from __future__ import annotations

import argparse
import logging
import re
import sys
import time
from dataclasses import dataclass, asdict, fields
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

START_URL = "https://books.toscrape.com/catalogue/page-1.html"
USER_AGENT = "Mozilla/5.0 (compatible; portfolio-scraper/1.0)"
RATING_WORDS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}

log = logging.getLogger("scraper")


@dataclass
class Book:
    title: str
    price_gbp: float
    rating: int
    in_stock: bool
    product_url: str
    image_url: str
    source_page: int


# ---------------------------------------------------------------- fetching
def make_session(retries: int = 3) -> requests.Session:
    from requests.adapters import HTTPAdapter
    from urllib3.util.retry import Retry

    s = requests.Session()
    s.headers["User-Agent"] = USER_AGENT
    retry = Retry(total=retries, backoff_factor=1.0, status_forcelist=(429, 500, 502, 503, 504))
    s.mount("https://", HTTPAdapter(max_retries=retry))
    s.mount("http://", HTTPAdapter(max_retries=retry))
    return s


def fetch(session: requests.Session, url: str) -> BeautifulSoup:
    resp = session.get(url, timeout=20)
    resp.raise_for_status()
    resp.encoding = "utf-8"
    return BeautifulSoup(resp.text, "html.parser")


# ---------------------------------------------------------------- parsing / cleaning
def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def parse_price(raw: str) -> float:
    digits = re.sub(r"[^\d.]", "", raw or "")
    return round(float(digits), 2) if digits else 0.0


def parse_page(soup: BeautifulSoup, page_url: str, page_no: int) -> list[Book]:
    books = []
    for card in soup.select("article.product_pod"):
        link = card.select_one("h3 a")
        rating_cls = card.select_one("p.star-rating")["class"]
        rating = next((RATING_WORDS[c] for c in rating_cls if c in RATING_WORDS), 0)
        books.append(
            Book(
                title=clean_text(link.get("title") or link.text),
                price_gbp=parse_price(card.select_one("p.price_color").text),
                rating=rating,
                in_stock="in stock" in clean_text(card.select_one("p.availability").text).lower(),
                product_url=urljoin(page_url, link["href"]),
                image_url=urljoin(page_url, card.select_one("img")["src"]),
                source_page=page_no,
            )
        )
    return books


def next_page_url(soup: BeautifulSoup, page_url: str) -> str | None:
    nxt = soup.select_one("li.next a")
    return urljoin(page_url, nxt["href"]) if nxt else None


def scrape(start_url: str, max_pages: int | None, delay: float) -> list[Book]:
    session = make_session()
    url, page_no, rows = start_url, 1, []
    while url and (max_pages is None or page_no <= max_pages):
        log.info("page %d: %s", page_no, url)
        soup = fetch(session, url)
        rows.extend(parse_page(soup, url, page_no))
        url = next_page_url(soup, url)
        page_no += 1
        if url and delay:
            time.sleep(delay)  # be polite
    return rows


def dedupe(rows: list[Book]) -> list[Book]:
    seen, out = set(), []
    for r in rows:
        if r.product_url not in seen:
            seen.add(r.product_url)
            out.append(r)
    return out


# ---------------------------------------------------------------- export
HEADERS = {
    "title": "Title",
    "price_gbp": "Price (GBP)",
    "rating": "Rating (1-5)",
    "in_stock": "In Stock",
    "product_url": "Product URL",
    "image_url": "Image URL",
    "source_page": "Source Page",
}


def export_xlsx(rows: list[Book], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Books"
    keys = [f.name for f in fields(Book)]
    ws.append([HEADERS[k] for k in keys])
    for r in rows:
        d = asdict(r)
        d["in_stock"] = "Yes" if d["in_stock"] else "No"
        ws.append([d[k] for k in keys])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    price_col = keys.index("price_gbp") + 1
    for (cell,) in ws.iter_rows(min_row=2, min_col=price_col, max_col=price_col):
        cell.number_format = "£#,##0.00"

    autofit(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # small summary sheet
    s = wb.create_sheet("Summary")
    prices = [r.price_gbp for r in rows] or [0]
    for label, value in [
        ("Total books", len(rows)),
        ("In stock", sum(r.in_stock for r in rows)),
        ("Average price (GBP)", round(sum(prices) / len(prices), 2)),
        ("Min price (GBP)", min(prices)),
        ("Max price (GBP)", max(prices)),
        ("5-star books", sum(r.rating == 5 for r in rows)),
    ]:
        s.append([label, value])
    for cell in s["A"]:
        cell.font = Font(bold=True)
    autofit(s)
    wb.save(path)


def autofit(ws, min_width: int = 8, max_width: int = 70) -> None:
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_width, min(width + 2, max_width))


# ---------------------------------------------------------------- CLI
def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--start-url", default=START_URL)
    ap.add_argument("--max-pages", type=int, default=None, help="stop after N pages (default: all)")
    ap.add_argument("--delay", type=float, default=0.5, help="seconds between requests")
    ap.add_argument("--out", type=Path, default=Path("output/books.xlsx"))
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    raw = scrape(args.start_url, args.max_pages, args.delay)
    rows = dedupe(raw)
    export_xlsx(rows, args.out)
    log.info("scraped %d rows, %d after de-duplication -> %s", len(raw), len(rows), args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
