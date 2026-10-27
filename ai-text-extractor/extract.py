"""Extract structured fields from unstructured text (invoices, e-mails) into Excel using an LLM.

    python extract.py samples/*.txt -o output/extracted.xlsx                 # LLM (OpenAI-compatible)
    python extract.py samples/*.txt -o output/extracted.xlsx --backend rules # offline, no API key

LLM settings come from environment variables (or a .env file next to this script):
    LLM_BASE_URL   default https://api.openai.com/v1   (e.g. http://localhost:11434/v1 for Ollama)
    LLM_API_KEY    your key (not needed for local servers)
    LLM_MODEL      default gpt-4o-mini

Every answer is validated: types converted, dates normalised, enums checked, and invoice
arithmetic verified (line items -> subtotal, subtotal + tax -> total). Anything suspicious
goes to a "Review" sheet instead of silently into the data.
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backends import get_backend
from schema import DOC_TYPES

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")


def load_dotenv(path: Path):
    if not path.exists():
        return
    for line in path.read_text().splitlines():
        if "=" in line and not line.lstrip().startswith("#"):
            k, v = line.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


# ---------------------------------------------------------------- normalisation + validation
def to_number(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".") if len(s.split(",")[-1]) == 2 else s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def to_date(v):
    if not v:
        return None
    if isinstance(v, date):
        return v
    for fmt in ("%Y-%m-%d", "%d %b %Y", "%d %B %Y", "%m/%d/%Y", "%d.%m.%Y", "%B %d, %Y"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            pass
    return None


def normalise(raw: dict) -> tuple[dict, list[str]]:
    """Convert types per schema.py and collect problems."""
    problems = []
    doc_type = str(raw.get("doc_type", "")).lower().strip()
    if doc_type not in DOC_TYPES:
        return {"doc_type": "unknown", "fields": {}, "line_items": []}, [f"unknown doc_type {raw.get('doc_type')!r}"]
    spec = DOC_TYPES[doc_type]

    def conv(name, typ, v):
        if v in (None, "", "null", "N/A"):
            return None
        if typ == "number":
            n = to_number(v)
            if n is None:
                problems.append(f"{name}: not a number ({v!r})")
            return n
        if typ == "date":
            d = to_date(v)
            if d is None:
                problems.append(f"{name}: unreadable date ({v!r})")
            return d
        if typ.startswith("enum:"):
            allowed = typ[5:].split("|")
            s = str(v).lower().strip()
            if s not in allowed:
                problems.append(f"{name}: {v!r} not one of {allowed}")
                return None
            return s
        return re.sub(r"\s+", " ", str(v)).strip()

    fields = {f: conv(f, typ, raw.get("fields", {}).get(f)) for f, (typ, _d) in spec["fields"].items()}
    items = []
    for i, it in enumerate(raw.get("line_items") or [], 1):
        items.append({f: conv(f"line {i} {f}", typ, it.get(f)) for f, (typ, _d) in spec.get("line_items", {}).items()})
    return {"doc_type": doc_type, "fields": fields, "line_items": items}, problems


def check(doc: dict) -> list[str]:
    """Business rules - the cheapest way to catch a wrong extraction."""
    p = []
    f, items = doc["fields"], doc["line_items"]
    if doc["doc_type"] == "invoice":
        for req in ("vendor", "invoice_number", "invoice_date", "total"):
            if f.get(req) in (None, ""):
                p.append(f"missing {req}")
        for i, it in enumerate(items, 1):
            if None not in (it["quantity"], it["unit_price"], it["amount"]) and abs(it["quantity"] * it["unit_price"] - it["amount"]) > 0.011:
                p.append(f"line {i}: {it['quantity']} x {it['unit_price']} != {it['amount']}")
        if items and f.get("subtotal") is not None and all(it["amount"] is not None for it in items):
            s = round(sum(it["amount"] for it in items), 2)
            if abs(s - f["subtotal"]) > 0.011:
                p.append(f"line items sum to {s}, subtotal says {f['subtotal']}")
        if None not in (f.get("subtotal"), f.get("tax"), f.get("total")) and abs(f["subtotal"] + f["tax"] - f["total"]) > 0.011:
            p.append(f"subtotal + tax = {round(f['subtotal'] + f['tax'], 2)}, total says {f['total']}")
        if f.get("due_date") and f.get("invoice_date") and f["due_date"] < f["invoice_date"]:
            p.append("due date before invoice date")
    elif doc["doc_type"] == "email":
        for req in ("sender_name", "intent"):
            if not f.get(req):
                p.append(f"missing {req}")
    return p


def process(paths: list[Path], backend) -> list[dict]:
    results = []
    for path in paths:
        text = path.read_text(encoding="utf-8", errors="replace")
        try:
            raw = backend.extract(text)
            doc, problems = normalise(raw)
            problems += check(doc)
        except Exception as e:                          # one bad file never stops the batch
            doc, problems = {"doc_type": "unknown", "fields": {}, "line_items": []}, [f"extraction failed: {e}"]
        doc.update(file=path.name, problems=problems, backend=backend.name)
        results.append(doc)
        status = "OK" if not problems else f"REVIEW ({len(problems)})"
        print(f"{path.name:<20} {doc['doc_type']:<8} {status}")
    return results


# ---------------------------------------------------------------- Excel
def _sheet(wb, title, headers, rows, review_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
        if review_col is not None and row[review_col].value:
            for cell in row:
                cell.fill = REVIEW_FILL
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(10, min(60, w + 2))
    return ws


def write_excel(results: list[dict], out: Path):
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    inv_fields = list(DOC_TYPES["invoice"]["fields"])
    mail_fields = list(DOC_TYPES["email"]["fields"])
    inv = [r for r in results if r["doc_type"] == "invoice"]
    mails = [r for r in results if r["doc_type"] == "email"]
    _sheet(wb, "Invoices", ["file"] + inv_fields + ["line_items", "needs_review"],
           [[r["file"]] + [r["fields"][f] for f in inv_fields] + [len(r["line_items"]), "yes" if r["problems"] else ""] for r in inv],
           review_col=len(inv_fields) + 2)
    li_fields = list(DOC_TYPES["invoice"]["line_items"])
    _sheet(wb, "Line Items", ["file", "invoice_number"] + li_fields,
           [[r["file"], r["fields"]["invoice_number"]] + [it[f] for f in li_fields] for r in inv for it in r["line_items"]])
    _sheet(wb, "Emails", ["file"] + mail_fields + ["needs_review"],
           [[r["file"]] + [r["fields"][f] for f in mail_fields] + ["yes" if r["problems"] else ""] for r in mails],
           review_col=len(mail_fields) + 1)
    _sheet(wb, "Review", ["file", "doc_type", "problem"],
           [[r["file"], r["doc_type"], p] for r in results for p in r["problems"]])
    wb.save(out)


def main(argv=None) -> int:
    load_dotenv(Path(__file__).with_name(".env"))
    ap = argparse.ArgumentParser(description="Unstructured text -> structured Excel with an LLM")
    ap.add_argument("inputs", nargs="+", help="text files (globs allowed)")
    ap.add_argument("-o", "--out", default="output/extracted.xlsx")
    ap.add_argument("--backend", choices=["openai", "rules"], default=os.environ.get("EXTRACT_BACKEND", "openai"))
    ap.add_argument("--model", help="override LLM_MODEL")
    ap.add_argument("--json", help="also save the normalised results as JSON")
    a = ap.parse_args(argv)
    paths = sorted({Path(p) for pat in a.inputs for p in (glob.glob(pat) or [pat])})
    missing = [p for p in paths if not p.exists()]
    if missing:
        print("not found:", ", ".join(map(str, missing)), file=sys.stderr)
        return 1
    if a.backend == "openai" and not os.environ.get("LLM_API_KEY") and not re.search(r"localhost|127\.0\.0\.1", os.environ.get("LLM_BASE_URL", "")):
        print("LLM_API_KEY is not set. Set it (see .env.example) or run with --backend rules.", file=sys.stderr)
        return 2
    backend = get_backend(a.backend, **({"model": a.model} if a.model and a.backend == "openai" else {}))
    results = process(paths, backend)
    write_excel(results, Path(a.out))
    if a.json:
        Path(a.json).parent.mkdir(parents=True, exist_ok=True)
        Path(a.json).write_text(json.dumps(results, indent=2, default=str) + "\n")
    n_rev = sum(1 for r in results if r["problems"])
    print(f"{len(results)} documents -> {a.out} ({n_rev} flagged for review)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
