"""Offline tests. No API key and no internet:
  * the rules backend is checked against hand-verified expected values (expected/*.json)
  * the OpenAI-compatible backend talks to a local fake server that imitates /v1/chat/completions
Run with:  python -m unittest -v
"""
import json
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

from openpyxl import load_workbook

import extract
from backends import OpenAICompatibleBackend, RuleBackend, build_prompt, parse_json_answer

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = sorted((ROOT / "samples").glob("*.txt"))


class RuleBackendTests(unittest.TestCase):
    def test_matches_expected_values(self):
        for path in SAMPLES:
            with self.subTest(path.name):
                doc, problems = extract.normalise(RuleBackend().extract(path.read_text()))
                expected = json.loads((ROOT / "expected" / (path.stem + ".json")).read_text())
                got = json.loads(json.dumps({k: doc[k] for k in ("doc_type", "fields", "line_items")}, default=str))
                self.assertEqual(got, expected)
                self.assertEqual(problems + extract.check(doc), [])


# ------------------------------------------------------------------ fake OpenAI-compatible server
class FakeLLM(BaseHTTPRequestHandler):
    replies = []          # list of (status, content) consumed in order
    requests = []

    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        FakeLLM.requests.append({"path": self.path, "auth": self.headers.get("Authorization"), "body": body})
        status, content = FakeLLM.replies.pop(0)
        payload = {"choices": [{"message": {"role": "assistant", "content": content}}]} if status == 200 else {"error": content}
        data = json.dumps(payload).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def log_message(self, *a):
        pass


INVOICE_JSON = {
    "doc_type": "invoice",
    "fields": {"vendor": "Sample Supplies Ltd", "customer": "Demo Bakery Co.", "invoice_number": "INV-2026-0147",
               "invoice_date": "2026-09-03", "due_date": "3 Oct 2026", "currency": "GBP",
               "subtotal": "258.00", "tax": 51.6, "total": "£309.60"},
    "line_items": [{"description": "Paper cups", "quantity": 4, "unit_price": 38, "amount": 152},
                   {"description": "Kraft bags", "quantity": 2, "unit_price": 24.5, "amount": 49},
                   {"description": "Napkins", "quantity": 3, "unit_price": 19, "amount": 57}],
}


class OpenAIBackendTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = HTTPServer(("127.0.0.1", 0), FakeLLM)
        threading.Thread(target=cls.server.serve_forever, daemon=True).start()
        cls.base = f"http://127.0.0.1:{cls.server.server_port}/v1"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()

    def setUp(self):
        FakeLLM.replies, FakeLLM.requests = [], []

    def backend(self, **kw):
        import requests
        s = requests.Session()
        s.trust_env = False                       # talk to the local fake server directly
        return OpenAICompatibleBackend(base_url=self.base, api_key="test-key", model="test-model",
                                       session=s, sleep=lambda s: None, **kw)

    def test_request_format_and_normalisation(self):
        FakeLLM.replies = [(200, json.dumps(INVOICE_JSON))]
        doc, problems = extract.normalise(self.backend().extract("some invoice text"))
        req = FakeLLM.requests[0]
        self.assertEqual(req["path"], "/v1/chat/completions")
        self.assertEqual(req["auth"], "Bearer test-key")
        self.assertEqual(req["body"]["model"], "test-model")
        self.assertEqual(req["body"]["temperature"], 0)
        self.assertEqual(req["body"]["response_format"], {"type": "json_object"})
        self.assertEqual(req["body"]["messages"][1]["content"], "some invoice text")
        self.assertIn("invoice_number", req["body"]["messages"][0]["content"])   # schema is in the prompt
        self.assertEqual(problems, [])
        self.assertEqual(doc["fields"]["total"], 309.6)                           # "£309.60" -> number
        self.assertEqual(str(doc["fields"]["due_date"]), "2026-10-03")            # "3 Oct 2026" -> date
        self.assertEqual(extract.check(doc), [])

    def test_fenced_json_and_retry_on_invalid_json(self):
        FakeLLM.replies = [(200, "Sure! Here is the data: {broken"),
                           (200, "```json\n" + json.dumps(INVOICE_JSON) + "\n```")]
        raw = self.backend().extract("text")
        self.assertEqual(raw["fields"]["invoice_number"], "INV-2026-0147")
        self.assertEqual(len(FakeLLM.requests), 2)
        self.assertIn("not valid JSON", FakeLLM.requests[1]["body"]["messages"][-1]["content"])

    def test_retries_server_errors_and_falls_back_without_json_mode(self):
        FakeLLM.replies = [(503, "busy"), (400, "response_format is not supported"), (200, json.dumps(INVOICE_JSON))]
        raw = self.backend().extract("text")
        self.assertEqual(raw["doc_type"], "invoice")
        self.assertNotIn("response_format", FakeLLM.requests[-1]["body"])

    def test_auth_error_is_not_retried(self):
        FakeLLM.replies = [(401, "invalid api key")]
        with self.assertRaisesRegex(RuntimeError, "401"):
            self.backend().extract("text")
        self.assertEqual(len(FakeLLM.requests), 1)

    def test_end_to_end_excel_with_fake_llm(self):
        bad = json.loads(json.dumps(INVOICE_JSON))
        bad["fields"]["total"] = 320.00                          # model misread the total
        FakeLLM.replies = [(200, json.dumps(INVOICE_JSON)), (200, json.dumps(bad)), (200, "not json"),
                           (200, "still not json"), (200, "nope")]
        with tempfile.TemporaryDirectory() as d:
            files = []
            for i in range(3):
                p = Path(d) / f"doc{i}.txt"; p.write_text("text"); files.append(p)
            results = extract.process(files, self.backend())
            out = Path(d) / "x.xlsx"
            extract.write_excel(results, out)
            wb = load_workbook(out)
            self.assertEqual(wb.sheetnames, ["Invoices", "Line Items", "Emails", "Review"])
            inv = list(wb["Invoices"].iter_rows(min_row=2, values_only=True))
            self.assertEqual([r[-1] for r in inv], [None, "yes"])        # 2nd invoice flagged
            review = [r[2] for r in wb["Review"].iter_rows(min_row=2, values_only=True)]
            self.assertTrue(any("total says 320.0" in r for r in review))
            self.assertTrue(any("extraction failed" in r for r in review))  # 3rd file: 3 bad answers
            self.assertEqual(len(list(wb["Line Items"].iter_rows(min_row=2))), 6)


class ValidationTests(unittest.TestCase):
    def test_parse_json_answer_variants(self):
        self.assertEqual(parse_json_answer('{"doc_type": "email"}')["line_items"], [])
        self.assertEqual(parse_json_answer('Result:\n{"doc_type": "email", "fields": {}}\nDone')["doc_type"], "email")
        with self.assertRaises(ValueError):
            parse_json_answer("no json here")

    def test_bad_enum_date_and_arithmetic_are_reported(self):
        doc, problems = extract.normalise({"doc_type": "email", "fields": {"sender_name": "A", "intent": "angry",
                                                                           "deadline": "someday"}})
        self.assertTrue(any(p.startswith("intent: 'angry' not one of") for p in problems))
        self.assertTrue(any("unreadable date" in p for p in problems))
        inv, _ = extract.normalise({"doc_type": "invoice", "fields": {"vendor": "V", "invoice_number": "1",
                                    "invoice_date": "2026-01-10", "due_date": "2026-01-01", "subtotal": 10, "tax": 2, "total": 13},
                                    "line_items": [{"description": "x", "quantity": 2, "unit_price": 4, "amount": 9}]})
        p = extract.check(inv)
        self.assertEqual(len(p), 4)          # line math, items vs subtotal, subtotal+tax, due before issue

    def test_prompt_lists_every_schema_field(self):
        system = build_prompt("x")[0]["content"]
        for f in ("vendor", "line_items", "urgency", "deadline"):
            self.assertIn(f, system)

    def test_cli_rules_backend(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "o.xlsx"
            self.assertEqual(extract.main([str(ROOT / "samples" / "*.txt"), "-o", str(out), "--backend", "rules"]), 0)
            wb = load_workbook(out)
            self.assertEqual(wb["Invoices"].max_row, 4)
            self.assertEqual(wb["Emails"].max_row, 4)
            self.assertEqual(wb["Line Items"].max_row, 10)
            self.assertEqual(wb["Review"].max_row, 1)


if __name__ == "__main__":
    unittest.main()
