#!/usr/bin/env python3
"""Extract a table that spans multiple PDF pages into one clean Excel sheet and CSV.

Handles:
- header row repeated on every page (kept once)
- continuation pages without a header (appended, not lost)
- page titles / footers outside the table (ignored)
- numeric cleanup ("1,234.50" -> 1234.5) and empty-row removal

Usage:
    python extract.py samples/sales_report.pdf
    python extract.py input.pdf --out-dir output --numeric "Qty,Unit Price,Total"
    python extract.py input.pdf --pages 2-5
"""
from __future__ import annotations

import argparse
import csv
import logging
import re
import sys
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger("extract")

# Ruled tables (cell borders drawn) -> "lines" strategy is the most reliable.
# For borderless tables switch both strategies to "text".
TABLE_SETTINGS = {"vertical_strategy": "lines", "horizontal_strategy": "lines", "snap_tolerance": 3}


def norm(cell) -> str:
    return re.sub(r"\s+", " ", str(cell or "")).strip()


def parse_pages(spec: str | None, total: int) -> list[int]:
    if not spec:
        return list(range(total))
    out = []
    for part in spec.split(","):
        a, _, b = part.partition("-")
        out.extend(range(int(a) - 1, int(b or a)))
    return [p for p in out if 0 <= p < total]


def extract_rows(pdf_path: Path, pages: str | None = None) -> tuple[list[str], list[list[str]], list[dict]]:
    header: list[str] | None = None
    rows: list[list[str]] = []
    report = []
    with pdfplumber.open(pdf_path) as pdf:
        for idx in parse_pages(pages, len(pdf.pages)):
            page = pdf.pages[idx]
            tables = page.extract_tables(TABLE_SETTINGS)
            page_rows = 0
            for table in tables:
                cleaned = [[norm(c) for c in r] for r in table]
                cleaned = [r for r in cleaned if any(r)]  # drop empty rows
                if not cleaned:
                    continue
                if header is None:
                    header, cleaned = cleaned[0], cleaned[1:]
                elif cleaned[0] == header:
                    cleaned = cleaned[1:]  # repeated header on a later page
                # continuation page: first row is data, keep everything
                for r in cleaned:
                    if len(r) != len(header):
                        log.warning("page %d: skipping row with %d cols (expected %d): %s", idx + 1, len(r), len(header), r)
                        continue
                    rows.append(r)
                    page_rows += 1
            report.append({"page": idx + 1, "tables": len(tables), "rows": page_rows})
            log.info("page %d: %d table(s), %d data rows", idx + 1, len(tables), page_rows)
    if header is None:
        raise SystemExit(f"No table found in {pdf_path}. Try TABLE_SETTINGS with the 'text' strategy.")
    return header, rows, report


def to_number(value: str):
    v = value.replace(",", "").replace("$", "").replace("£", "").replace("€", "").strip()
    if re.fullmatch(r"-?\d+", v):
        return int(v)
    if re.fullmatch(r"-?\d+\.\d+", v):
        return float(v)
    return value


def detect_numeric(header: list[str], rows: list[list[str]]) -> list[str]:
    cols = []
    for i, h in enumerate(header):
        vals = [r[i] for r in rows if r[i]]
        if vals and all(isinstance(to_number(v), (int, float)) for v in vals):
            cols.append(h)
    return cols


def write_csv(header, rows, path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig opens cleanly in Excel
        csv.writer(f).writerows([header, *rows])


def write_xlsx(header, rows, numeric_cols, report, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    num_idx = {header.index(c) for c in numeric_cols if c in header}
    ws.append(header)
    for r in rows:
        ws.append([to_number(v) if i in num_idx else v for i, v in enumerate(r)])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for i in num_idx:
        for (cell,) in ws.iter_rows(min_row=2, min_col=i + 1, max_col=i + 1):
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
    autofit(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("Extraction report")
    s.append(["Page", "Tables found", "Data rows"])
    for r in report:
        s.append([r["page"], r["tables"], r["rows"]])
    s.append(["Total", sum(r["tables"] for r in report), sum(r["rows"] for r in report)])
    for cell in s[1]:
        cell.font = Font(bold=True)
    autofit(s)
    wb.save(path)


def autofit(ws, min_width=8, max_width=60):
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_width, min(width + 2, max_width))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf", type=Path)
    ap.add_argument("--out-dir", type=Path, default=Path("output"))
    ap.add_argument("--pages", help="e.g. 1-3,5 (default: all pages)")
    ap.add_argument("--numeric", help="comma-separated columns to convert to numbers (default: auto-detect)")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")

    header, rows, report = extract_rows(args.pdf, args.pages)
    numeric = [c.strip() for c in args.numeric.split(",")] if args.numeric else detect_numeric(header, rows)
    args.out_dir.mkdir(parents=True, exist_ok=True)
    stem = args.pdf.stem
    write_csv(header, rows, args.out_dir / f"{stem}.csv")
    write_xlsx(header, rows, numeric, report, args.out_dir / f"{stem}.xlsx")
    log.info("%d rows x %d columns -> %s/%s.xlsx and .csv (numeric: %s)",
             len(rows), len(header), args.out_dir, stem, ", ".join(numeric) or "none")
    return 0


if __name__ == "__main__":
    sys.exit(main())
