"""Round-trip test: generate a PDF, extract it, compare with the source data.
Run with:  python -m unittest
"""
import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import extract
import make_sample_pdf as sample


class RoundTripTest(unittest.TestCase):
    def test_multi_page_table_is_merged_exactly(self):
        rows = sample.make_rows(100, seed=7)
        with tempfile.TemporaryDirectory() as d:
            pdf = Path(d) / "t.pdf"
            n_pages = sample.build_pdf(rows, pdf)
            self.assertGreaterEqual(n_pages, 3)  # includes the header-less page 3

            header, got, report = extract.extract_rows(pdf)
            self.assertEqual(header, sample.HEADER)
            self.assertEqual(got, rows)                       # no lost, duplicated or header rows
            self.assertEqual(sum(r["rows"] for r in report), 100)

            out = Path(d) / "out"
            self.assertEqual(extract.main([str(pdf), "--out-dir", str(out)]), 0)
            with (out / "t.csv").open(encoding="utf-8-sig") as f:
                self.assertEqual(list(csv.reader(f)), [sample.HEADER, *rows])
            ws = load_workbook(out / "t.xlsx")["Data"]
            self.assertTrue(ws["A1"].font.bold)
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertIsInstance(ws["H2"].value, float)      # "1,234.50" became a number

    def test_number_cleanup(self):
        self.assertEqual(extract.to_number("1,234.50"), 1234.5)
        self.assertEqual(extract.to_number("$12"), 12)
        self.assertEqual(extract.to_number("North"), "North")


if __name__ == "__main__":
    unittest.main()
