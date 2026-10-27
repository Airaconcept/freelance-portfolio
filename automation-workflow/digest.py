#!/usr/bin/env python3
"""Config-driven monitor: RSS feeds + web pages -> new items only -> daily Markdown/Excel report
-> optional webhook push (Slack, Discord or generic JSON).

Usage:
    python digest.py                       # uses config.yaml
    python digest.py --config my.yaml
    python digest.py --dry-run             # build the report, print the webhook payload, send nothing
    python digest.py --reset               # forget what was seen and start over

Schedule it with cron (Linux/macOS) or Task Scheduler (Windows), e.g. every day at 08:00:
    0 8 * * *  cd /path/to/automation-workflow && python digest.py
"""
from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import sys
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path

import feedparser
import requests
import yaml
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

USER_AGENT = "Mozilla/5.0 (compatible; daily-digest/1.0)"
log = logging.getLogger("digest")


@dataclass
class Item:
    source: str
    title: str
    link: str
    published: str
    uid: str


# ---------------------------------------------------------------- collectors
def fetch_url(url: str) -> bytes:
    r = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=20)
    r.raise_for_status()
    return r.content


def collect_rss(src: dict, limit: int) -> list[Item]:
    feed = feedparser.parse(fetch_url(src["url"]))
    items = []
    for e in feed.entries[:limit]:
        link = e.get("link", "")
        items.append(Item(
            source=src["name"],
            title=" ".join(e.get("title", "(no title)").split()),
            link=link,
            published=e.get("published", e.get("updated", "")),
            uid=e.get("id") or link or e.get("title", ""),
        ))
    return items


def collect_web(src: dict, limit: int) -> list[Item]:
    soup = BeautifulSoup(fetch_url(src["url"]), "html.parser")
    items = []
    for el in soup.select(src["selector"])[:limit]:
        text = " ".join(el.get_text(" ").split())
        if not text:
            continue
        a = el if el.name == "a" else el.find("a")
        link = requests.compat.urljoin(src["url"], a["href"]) if a and a.get("href") else src["url"]
        items.append(Item(src["name"], text[:200], link, "",
                          hashlib.sha1(f'{src["url"]}|{text}'.encode()).hexdigest()))
    return items


COLLECTORS = {"rss": collect_rss, "web": collect_web}


# ---------------------------------------------------------------- state
def load_state(path: Path) -> dict:
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {"seen": {}, "runs": 0}


def save_state(path: Path, state: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(state, indent=2), encoding="utf-8")


def keyword_match(title: str, keywords: list[str]) -> bool:
    return not keywords or any(k.lower() in title.lower() for k in keywords)


# ---------------------------------------------------------------- reports
def render_markdown(title: str, now: datetime, new: list[Item], errors: list[str], totals: dict) -> str:
    lines = [f"# {title} - {now:%Y-%m-%d}", "",
             f"Generated {now:%Y-%m-%d %H:%M %Z}. **{len(new)} new item(s)** from {len(totals)} source(s).", ""]
    for source, count in totals.items():
        group = [i for i in new if i.source == source]
        lines.append(f"## {source} ({len(group)} new / {count} checked)")
        lines.append("")
        lines += [f"- [{i.title}]({i.link})" for i in group] or ["_No new items._"]
        lines.append("")
    if errors:
        lines += ["## Errors", ""] + [f"- {e}" for e in errors] + [""]
    return "\n".join(lines)


def write_excel(path: Path, new: list[Item]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "New items"
    ws.append(["Source", "Title", "Link", "Published"])
    for i in new:
        ws.append([i.source, i.title, i.link, i.published])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(10, min(w + 2, 80))
    ws.freeze_panes = "A2"
    wb.save(path)


def webhook_payload(fmt: str, markdown: str, new: list[Item], title: str) -> dict:
    short = markdown if len(markdown) < 1900 else markdown[:1900] + "\n..."
    if fmt == "slack":
        return {"text": short}
    if fmt == "discord":
        return {"content": short}
    return {"title": title, "count": len(new), "items": [asdict(i) for i in new], "markdown": markdown}


# ---------------------------------------------------------------- main
def run(cfg: dict, base: Path, dry_run: bool = False, now: datetime | None = None) -> dict:
    now = now or datetime.now(timezone.utc).astimezone()
    rep = cfg.get("report", {})
    limit = int(rep.get("max_items_per_source", 15))
    state_path = base / rep.get("state_file", "state/seen.json")
    out_dir = base / rep.get("output_dir", "reports")
    state = load_state(state_path)
    first_run = state["runs"] == 0
    keywords = cfg.get("keywords") or []

    new, errors, totals = [], [], {}
    for src in cfg.get("sources", []):
        try:
            items = COLLECTORS[src["type"]](src, limit)
        except Exception as exc:  # one broken source must not kill the report
            errors.append(f'{src.get("name")}: {exc.__class__.__name__}: {exc}')
            log.error(errors[-1])
            continue
        totals[src["name"]] = len(items)
        seen = set(state["seen"].get(src["name"], []))
        fresh = [i for i in items if i.uid not in seen and keyword_match(i.title, keywords)]
        if first_run and not rep.get("first_run_report_all", True):
            fresh = []
        new.extend(fresh)
        state["seen"][src["name"]] = list(seen | {i.uid for i in items})[-2000:]
        log.info("%s: %d checked, %d new", src["name"], len(items), len(fresh))

    title = rep.get("title", "Daily Digest")
    md = render_markdown(title, now, new, errors, totals)
    out_dir.mkdir(parents=True, exist_ok=True)
    md_path = out_dir / f"digest-{now:%Y-%m-%d_%H%M}.md"
    xlsx_path = out_dir / f"digest-{now:%Y-%m-%d_%H%M}.xlsx"
    md_path.write_text(md, encoding="utf-8")
    write_excel(xlsx_path, new)

    hook = cfg.get("webhook", {}) or {}
    url = os.environ.get("DIGEST_WEBHOOK_URL") or hook.get("url")
    sent = False
    if hook.get("enabled") and url:
        payload = webhook_payload(hook.get("format", "slack"), md, new, title)
        if dry_run:
            print(json.dumps(payload, indent=2)[:3000])
        else:
            r = requests.post(url, json=payload, timeout=20)
            r.raise_for_status()
            sent = True
            log.info("webhook delivered (%s)", r.status_code)

    if not dry_run:
        state["runs"] += 1
        state["last_run"] = now.isoformat()
        save_state(state_path, state)
    return {"new": len(new), "errors": errors, "markdown": str(md_path), "excel": str(xlsx_path), "webhook_sent": sent}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", type=Path, default=Path("config.yaml"))
    ap.add_argument("--dry-run", action="store_true", help="do not send the webhook or update state")
    ap.add_argument("--reset", action="store_true", help="delete saved state before running")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    cfg = yaml.safe_load(args.config.read_text(encoding="utf-8"))
    base = args.config.resolve().parent
    if args.reset:
        (base / cfg.get("report", {}).get("state_file", "state/seen.json")).unlink(missing_ok=True)
    result = run(cfg, base, dry_run=args.dry_run)
    log.info("done: %s", result)
    return 1 if result["errors"] and not result["new"] else 0


if __name__ == "__main__":
    sys.exit(main())
